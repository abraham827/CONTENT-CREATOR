from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document

from .generator import PostIdea


EXPORT_COLUMNS = ["Header", "Body", "Call To Action"]


def export_to_excel(ideas: Iterable[PostIdea], path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Posts"
    ws.append(EXPORT_COLUMNS)
    for idea in ideas:
        ws.append([idea.header, idea.body, idea.cta])
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 40)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def export_to_word(ideas: Iterable[PostIdea], path: Path) -> Path:
    doc = Document()
    doc.add_heading("Content Calendar Ideas", level=1)
    for index, idea in enumerate(ideas, start=1):
        doc.add_heading(f"Post {index}", level=2)
        doc.add_paragraph(f"Header: {idea.header}")
        doc.add_paragraph(f"Body: {idea.body}")
        doc.add_paragraph(f"Call to Action: {idea.cta}")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def export_to_powerpoint(ideas: Iterable[PostIdea], path: Path) -> Path:
    pres = Presentation()
    blank_slide_layout = pres.slide_layouts[6]
    for idea in ideas:
        slide = pres.slides.add_slide(blank_slide_layout)
        left = Inches(1)
        top = Inches(1)
        width = Inches(8)
        height = Inches(5)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.text = idea.header
        body_run = tf.add_paragraph().add_run()
        body_run.text = idea.body
        body_run.font.size = Pt(20)
        cta_run = tf.add_paragraph().add_run()
        cta_run.text = f"Call to Action: {idea.cta}"
        cta_run.font.size = Pt(18)
    path.parent.mkdir(parents=True, exist_ok=True)
    pres.save(path)
    return path
